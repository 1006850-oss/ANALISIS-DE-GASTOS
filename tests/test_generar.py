"""Pruebas del Hilo 7 (Excel de resultados, tablero e informe ejecutivo).

Corre la cadena completa (Hilos 1 a 6) sobre data sintética y luego genera las salidas.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "scripts"))
sys.path.insert(0, str(RAIZ / "tests"))

import alertas as al  # noqa: E402
import analizar as an  # noqa: E402
import clasificar as cl  # noqa: E402
import generar_excel as ge  # noqa: E402
import generar_informe as gi  # noqa: E402
import kraljic as kr  # noqa: E402
import limpiar_validar as lv  # noqa: E402
import proyectar as pr  # noqa: E402
from generar_sintetico import generar  # noqa: E402

PARAMS = lv.cargar_parametros()
RUC = re.compile(r"(?<!\d)(10|15|17|20)\d{9}(?!\d)")


@pytest.fixture(scope="module")
def salidas(tmp_path_factory):
    d = tmp_path_factory.mktemp("h7")
    df, _ = generar(filas=8000, semilla=71)
    t, _ = cl.clasificar(lv.procesar(df, PARAMS).tabla, cl.cargar_reglas())
    tabla = d / "tabla_clasificada.parquet"
    t.to_parquet(tabla, index=False)
    an.main(["--tabla", str(tabla), "--salida", str(d / "desc"), "--sin-graficos"])
    al.main(["--tabla", str(tabla), "--salida", str(d / "alert")])
    kr.main(["--tabla", str(tabla), "--salida", str(d / "kr")])
    pr.main(["--tabla", str(tabla), "--kraljic", str(d / "kr" / "tablas" / "kraljic_unidades.parquet"),
             "--salida", str(d / "plan")])
    comunes = ["--descriptivo", str(d / "desc"), "--alertas", str(d / "alert"), "--kraljic", str(d / "kr"),
               "--plan", str(d / "plan"), "--periodo", "prueba"]
    ge.main(comunes + ["--tabla", str(tabla), "--salida", str(d / "out")])
    codigo = gi.main(comunes + ["--salida", str(d / "out")])
    return {"d": d, "tabla": t, "codigo_informe": codigo, "argv": comunes}


def _textos(ruta: Path) -> list[str]:
    wb = load_workbook(ruta, read_only=True)
    return [str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v is not None]


# 1. Los totales del Excel y del informe cuadran con las tablas de origen.
def test_totales_cuadran(salidas):
    d = salidas["d"]
    T = ge.cargar_tablas({"desc": d / "desc", "alert": d / "alert", "kraljic": d / "kr", "plan": d / "plan"})
    control = T["desc.control"]
    gasto = control.loc[control["control"].str.startswith("Gasto analizado"), "valor_soles"].iloc[0]
    assert T["desc.pareto_proveedores"]["gasto"].sum() == pytest.approx(gasto)
    assert T["desc.gasto_categoria"]["gasto"].sum() == pytest.approx(gasto)
    wb = load_workbook(d / "out" / "resultados_prueba.xlsx", read_only=True)
    pareto = pd.DataFrame(list(wb["Pareto_ABC"].iter_rows(min_row=4, values_only=True)))
    pareto.columns = pareto.iloc[0]
    assert pd.to_numeric(pareto["gasto"].iloc[1:]).sum() == pytest.approx(gasto)
    c = gi.calcular_cifras(T)
    assert c.d["gasto"]["valor"] == pytest.approx(gasto)
    assert c.d["gasto"]["texto"] in "\n".join(_textos(d / "out" / "informe_ejecutivo_prueba.xlsx"))


# 2. Modelo en estrella: sin claves huérfanas y el total de hechos = gasto de la tabla clasificada.
def test_modelo_estrella_sin_huerfanos(salidas):
    carpeta = salidas["d"] / "out" / "tablero"
    h = pd.read_parquet(carpeta / "hechos_gasto.parquet")
    assert h["monto_gasto"].sum() == pytest.approx(salidas["tabla"]["monto_gasto"].sum())
    pares = [("dim_proveedor", "proveedor_id"), ("dim_sede", "sede"), ("dim_departamento", "departamento"),
             ("dim_comprador", "comprador_codigo"), ("dim_comprador", "aprobador_codigo")]
    for dim, col in pares:
        clave = "comprador_codigo" if dim == "dim_comprador" else col
        assert h[col].dropna().isin(pd.read_parquet(carpeta / f"{dim}.parquet")[clave]).all(), col
    assert h["fecha"].dropna().isin(pd.read_parquet(carpeta / "dim_fecha.parquet")["fecha"]).all()
    cat = pd.read_parquet(carpeta / "dim_categoria.parquet")
    assert not cat.duplicated().any()
    assert len(h.merge(cat, on=["categoria", "subcategoria", "producto"], how="left")) == len(h)


# 3. La verificación detecta una cifra alterada.
def test_verificacion_detecta_cifra_alterada(salidas, tmp_path):
    d = salidas["d"]
    assert salidas["codigo_informe"] == 0
    T = ge.cargar_tablas({"desc": d / "desc", "alert": d / "alert", "kraljic": d / "kr", "plan": d / "plan"})
    c = gi.calcular_cifras(T)
    ruta = d / "out" / "informe_ejecutivo_prueba.xlsx"
    assert gi.verificar_informe(ruta, c) == []
    wb = load_workbook(ruta)
    objetivo = c["gasto"]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cel in row:
                if isinstance(cel.value, str) and objetivo in cel.value:
                    cel.value = cel.value.replace(objetivo, "S/ 999.9 M")
    alterado = tmp_path / "alterado.xlsx"
    wb.save(alterado)
    problemas = gi.verificar_informe(alterado, c)
    assert any("'gasto'" in p for p in problemas)


# 4. Sin nombres de personas (compradores/aprobadores) ni RUC en el informe; sin nombres de personas en resultados.
def test_sin_nombres_ni_ruc(salidas):
    d = salidas["d"] / "out"
    informe = _textos(d / "informe_ejecutivo_prueba.xlsx")
    resultados = _textos(d / "resultados_prueba.xlsx")
    assert not any("PERSONA" in x for x in informe + resultados)
    assert not any(RUC.search(x) for x in informe)
    h = pd.read_parquet(d / "tablero" / "hechos_gasto.parquet")
    assert not h["comprador_codigo"].dropna().str.contains("PERSONA").any()


# 5. Máximo 5 hojas, cada una ajustada a una página.
def test_informe_cinco_hojas(salidas):
    wb = load_workbook(salidas["d"] / "out" / "informe_ejecutivo_prueba.xlsx")
    assert 1 <= len(wb.worksheets) <= 5
    for ws in wb.worksheets:
        assert ws.page_setup.fitToHeight == 1 and ws.page_setup.fitToWidth == 1
    assert "BORRADOR" in "\n".join(_textos(salidas["d"] / "out" / "informe_ejecutivo_prueba.xlsx"))


# 6. Reproducible: las cifras no cambian al volver a generar.
def test_reproducible(salidas, tmp_path):
    gi.main(salidas["argv"] + ["--salida", str(tmp_path)])
    a = pd.read_excel(salidas["d"] / "out" / "informe_trazabilidad.xlsx")
    b = pd.read_excel(tmp_path / "informe_trazabilidad.xlsx")
    pd.testing.assert_frame_equal(a, b)
