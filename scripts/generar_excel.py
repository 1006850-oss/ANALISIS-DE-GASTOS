"""Hilo 7 – Excel de resultados con tablero (sin Power BI).

Lee las tablas que ya produjeron los Hilos 1 a 6 (no recalcula indicadores) y arma un solo libro:
  Portada · Tablero_Resumen · Tablero_Proveedores · Tablero_Control · Tablero_Kraljic · Tablero_Plan ·
  Pareto_ABC · Frecuencia · Gasto_categoria · Gasto_subcategoria · Cruce_comprador_categoria · Concentracion_comprador ·
  Sedes · Alertas_resumen · Kraljic_unidades · Plan_anual · Calendario · Control_calidad · Glosario
También exporta el modelo en estrella (hechos + dimensiones) a <salida>/tablero/ para un futuro Power BI.

Uso:
  python scripts/generar_excel.py --descriptivo <dir Hilo 3> --alertas <dir Hilo 4> --kraljic <dir Hilo 5> \
         --plan <dir Hilo 6> --tabla <tabla_clasificada.parquet> [--calidad reporte_calidad.xlsx] \
         --periodo 2015-2017 --salida <carpeta>
"""
from __future__ import annotations

import argparse
import sys
from datetime import date
from pathlib import Path

import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parents[1]

AZUL, GRIS, TEXTO2 = "2A78D6", "F0EFEC", "52514E"
TITULO = Font(size=14, bold=True)
SUBT = Font(size=10, italic=True, color=TEXTO2)
CAB = Font(bold=True, color="FFFFFF")
CAB_FILL = PatternFill("solid", fgColor=AZUL)
KPI_FILL = PatternFill("solid", fgColor=GRIS)
SOLES = '"S/ "#,##0'
PCT = '0.0" %"'
ENTERO = "#,##0"


# --------------------------------------------------------------------------- carga

def cargar_tablas(dirs: dict[str, Path]) -> dict[str, pd.DataFrame]:
    t = {}
    for nombre, d in dirs.items():
        if d is None:
            continue
        for f in sorted((Path(d) / "tablas").glob("*.parquet")):
            t[f"{nombre}.{f.stem}"] = pd.read_parquet(f)
    return t


def leer_calidad(ruta: Path | None) -> pd.DataFrame:
    """Solo la sección de validaciones de reporte_calidad.xlsx (la hoja de códigos de personas no se lee)."""
    if ruta is None or not Path(ruta).exists():
        return pd.DataFrame({"nota": ["reporte_calidad.xlsx no disponible"]})
    crudo = pd.read_excel(ruta, sheet_name="Resumen", header=None)
    fila = crudo.index[crudo[0].astype(str) == "n"][0]
    v = crudo.iloc[fila + 1:, :6]
    v.columns = crudo.iloc[fila, :6].tolist()
    return v.dropna(how="all").reset_index(drop=True)


# --------------------------------------------------------------------------- utilidades de formato

def escribir_tabla(ws, df: pd.DataFrame, fila: int = 1, col: int = 1, soles=(), pct=(), enteros=(),
                   filtro: bool = True, ancho_max: int = 45) -> tuple[int, int]:
    """Escribe df con encabezado formateado; devuelve (primera fila de datos, última fila)."""
    for j, c in enumerate(df.columns):
        cel = ws.cell(row=fila, column=col + j, value=str(c))
        cel.font, cel.fill = CAB, CAB_FILL
        cel.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            if isinstance(v, float) and pd.isna(v):
                v = None
            elif hasattr(v, "item"):
                v = v.item()
            elif isinstance(v, pd.Timestamp):
                v = v.date()
            cel = ws.cell(row=fila + i, column=col + j, value=v)
            nombre = df.columns[j]
            if nombre in soles:
                cel.number_format = SOLES
            elif nombre in pct:
                cel.number_format = PCT
            elif nombre in enteros:
                cel.number_format = ENTERO
    for j, c in enumerate(df.columns):
        letra = get_column_letter(col + j)
        largo = max([len(str(c))] + [len(str(x)) for x in df.iloc[:200, j].tolist()])
        ws.column_dimensions[letra].width = max(ws.column_dimensions[letra].width or 0, min(ancho_max, largo + 2))
    if filtro and len(df):
        ws.auto_filter.ref = f"{get_column_letter(col)}{fila}:{get_column_letter(col + len(df.columns) - 1)}{fila + len(df)}"
    ws.freeze_panes = ws.cell(row=fila + 1, column=col)
    return fila + 1, fila + len(df)


def hoja_datos(wb, nombre: str, df: pd.DataFrame, fuente: str, **fmt) -> None:
    ws = wb.create_sheet(nombre[:31])
    ws["A1"] = nombre.replace("_", " ")
    ws["A1"].font = TITULO
    ws["A2"] = f"Fuente: {fuente}"
    ws["A2"].font = SUBT
    escribir_tabla(ws, df, fila=4, **fmt)


def kpi(ws, fila: int, col: int, etiqueta: str, valor, formato: str) -> None:
    a = ws.cell(row=fila, column=col, value=etiqueta)
    b = ws.cell(row=fila + 1, column=col, value=valor)
    a.font = Font(size=9, color=TEXTO2)
    b.font = Font(size=16, bold=True)
    b.number_format = formato
    for c in (a, b):
        c.fill = KPI_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(col)].width = 24


def grafico_barras(ws, datos_ws, col_cat: int, col_val: int, fila_ini: int, fila_fin: int, titulo: str,
                   ancla: str, horizontal: bool = False, alto: float = 7.5, ancho: float = 16) -> None:
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.title = titulo
    ch.legend = None
    ch.height, ch.width = alto, ancho
    ch.add_data(Reference(datos_ws, min_col=col_val, min_row=fila_ini - 1, max_row=fila_fin), titles_from_data=True)
    ch.set_categories(Reference(datos_ws, min_col=col_cat, min_row=fila_ini, max_row=fila_fin))
    ch.series[0].graphicalProperties.solidFill = AZUL
    ch.series[0].graphicalProperties.line.noFill = True
    if horizontal:
        ch.y_axis.scaling.orientation = "minMax"
        ch.x_axis.scaling.orientation = "maxMin"
    ws.add_chart(ch, ancla)


# --------------------------------------------------------------------------- libro

def _tablero_plan(tpl, T: dict) -> None:
    tpl["A1"] = "Tablero – Plan de compras (línea base, ejercicio metodológico)"
    tpl["A1"].font = TITULO
    mens = T["plan.plan_mensual"].groupby(["mes", "mes_nombre"])["linea_base"].sum().reset_index()
    mens = mens[["mes_nombre", "linea_base"]].rename(columns={"mes_nombre": "Mes", "linea_base": "Línea base"})
    ini, fin = escribir_tabla(tpl, mens, fila=3, soles=("Línea base",), filtro=False)
    ch = LineChart()
    ch.title, ch.legend, ch.height, ch.width = "Línea base mensual (recurrentes + variables, S/)", None, 7.5, 16
    ch.add_data(Reference(tpl, min_col=2, min_row=ini - 1, max_row=fin), titles_from_data=True)
    ch.set_categories(Reference(tpl, min_col=1, min_row=ini, max_row=fin))
    ch.series[0].graphicalProperties.line.solidFill = AZUL
    ch.series[0].graphicalProperties.line.width = 25000
    tpl.add_chart(ch, "D3")
    seg = T["plan.plan_plan"].groupby("segmento").agg(Unidades=("unidad", "size"), Linea_base=("linea_base", "sum")).reset_index()
    escribir_tabla(tpl, seg.rename(columns={"segmento": "Segmento", "Linea_base": "Línea base"}), fila=19,
                   soles=("Línea base",), enteros=("Unidades",), filtro=False)


def generar_resultados(T: dict[str, pd.DataFrame], calidad: pd.DataFrame, periodo: str, ruta: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Portada"
    reglas = yaml.safe_load(open(RAIZ / "config" / "reglas_taxonomia.yaml", encoding="utf-8"))["version"]
    alertas_v = yaml.safe_load(open(RAIZ / "config" / "reglas_alertas.yaml", encoding="utf-8"))["version"]
    ctrl = T["desc.control"].set_index("control")["valor_soles"]
    gasto = float([v for k, v in ctrl.items() if k.startswith("Gasto analizado")][0])
    filas = [("Análisis de gastos de compras – Resultados", ""), ("", ""),
             ("Periodo analizado", periodo), ("Gasto analizado (soles nominales)", gasto),
             ("Fecha de generación", date.today().isoformat()), ("Versión de reglas de taxonomía", reglas),
             ("Versión de reglas de alertas", alertas_v),
             ("Cómo leer", "Las hojas 'Tablero_…' resumen; las demás tienen el detalle con filtros. Las alertas son señales "
                           "para revisar, no conclusiones. Compradores y aprobadores con código (P001…)."),
             ("Advertencias", "Cuadrantes de Kraljic provisionales hasta el taller de expertos. Plan de compras = ejercicio "
                              "metodológico (data hasta 2017).")]
    for i, (a, b) in enumerate(filas, start=1):
        ws.cell(row=i, column=1, value=a).font = TITULO if i == 1 else Font(bold=True)
        c = ws.cell(row=i, column=2, value=b)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if isinstance(b, float):
            c.number_format = SOLES
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 36, 100

    # ---------------- Tablero Resumen
    tr = wb.create_sheet("Tablero_Resumen")
    tr["A1"] = "Tablero – Resumen"
    tr["A1"].font = TITULO
    tr["A2"] = f"Periodo {periodo} · soles nominales"
    tr["A2"].font = SUBT
    abc = T["desc.resumen_abc"].set_index("clase_abc")
    anio = T["desc.tendencia_categoria_anio"].groupby("anio_gasto")["gasto"].sum()
    alert = T["alert.resumen"]
    kpi(tr, 4, 1, "Gasto analizado", gasto, SOLES)
    kpi(tr, 4, 2, "Proveedores", int(len(T["desc.pareto_proveedores"])), ENTERO)
    kpi(tr, 4, 3, "Proveedores clase A", int(abc.loc["A", "proveedores"]), ENTERO)
    kpi(tr, 4, 4, "% gasto en clase A", float(abc.loc["A", "pct_gasto"]), PCT)
    kpi(tr, 4, 5, "Alertas de prioridad alta", int(alert["alta"].sum()), ENTERO)
    df_anio = pd.DataFrame({"Año": [str(a) for a in anio.index], "Gasto": anio.values})
    ini, fin = escribir_tabla(tr, df_anio, fila=8, col=1, soles=("Gasto",), filtro=False)
    grafico_barras(tr, tr, 1, 2, ini, fin, "Gasto por año de factura (S/)", "D8")
    cat = T["desc.gasto_categoria"][["categoria", "gasto", "pct_gasto"]].rename(
        columns={"categoria": "Categoría", "gasto": "Gasto", "pct_gasto": "% gasto"})
    ini, fin = escribir_tabla(tr, cat, fila=25, col=1, soles=("Gasto",), pct=("% gasto",), filtro=False)
    grafico_barras(tr, tr, 1, 2, ini, fin, "Gasto por categoría (S/)", "E25", horizontal=True, alto=10, ancho=18)

    # ---------------- Tablero Proveedores
    tp = wb.create_sheet("Tablero_Proveedores")
    tp["A1"] = "Tablero – Proveedores"
    tp["A1"].font = TITULO
    top = T["desc.pareto_proveedores"].head(20)[["ranking", "proveedor", "gasto", "pct_gasto", "pct_acumulado", "clase_abc"]]
    top = top.rename(columns={"ranking": "N°", "proveedor": "Proveedor", "gasto": "Gasto", "pct_gasto": "% gasto",
                              "pct_acumulado": "% acumulado", "clase_abc": "Clase"})
    ini, fin = escribir_tabla(tp, top, fila=3, soles=("Gasto",), pct=("% gasto", "% acumulado"), filtro=False)
    grafico_barras(tp, tp, 2, 3, ini, fin, "20 proveedores con más gasto (S/)", "H3", horizontal=True, alto=11, ancho=18)
    r_abc = T["desc.resumen_abc"].rename(columns={"clase_abc": "Clase", "proveedores": "Proveedores", "gasto": "Gasto",
                                                   "pct_proveedores": "% proveedores", "pct_gasto": "% gasto"})
    escribir_tabla(tp, r_abc[["Clase", "Proveedores", "% proveedores", "Gasto", "% gasto"]], fila=26,
                   soles=("Gasto",), pct=("% proveedores", "% gasto"), enteros=("Proveedores",), filtro=False)

    # ---------------- Tablero Control
    tc = wb.create_sheet("Tablero_Control")
    tc["A1"] = "Tablero – Control (señales para revisar, no conclusiones)"
    tc["A1"].font = TITULO
    a = alert.assign(tipo=lambda d: d["alerta"] + d["subtipo"].where(d["subtipo"] == "", " – " + d["subtipo"]))
    a = a[["tipo", "casos", "alta", "media", "baja", "monto_soles"]].rename(
        columns={"tipo": "Alerta", "casos": "Casos", "alta": "Alta", "media": "Media", "baja": "Baja", "monto_soles": "Monto"})
    ini, fin = escribir_tabla(tc, a, fila=3, soles=("Monto",), enteros=("Casos", "Alta", "Media", "Baja"), filtro=False)
    grafico_barras(tc, tc, 1, 3, ini, fin, "Casos de prioridad alta por alerta", "H3", horizontal=True, alto=10, ancho=18)

    # ---------------- Tablero Kraljic
    tk = wb.create_sheet("Tablero_Kraljic")
    tk["A1"] = "Tablero – Matriz de Kraljic (provisional hasta el taller de expertos)"
    tk["A1"].font = TITULO
    kr = T["kraljic.kraljic_resumen"].rename(columns={"cuadrante": "Cuadrante", "unidades": "Unidades", "gasto": "Gasto",
                                                     "pct_gasto": "% gasto"})
    escribir_tabla(tk, kr, fila=3, soles=("Gasto",), pct=("% gasto",), enteros=("Unidades",), filtro=False)
    png = Path(T.get("_kraljic_png", ""))
    if png.exists():
        from openpyxl.drawing.image import Image
        img = Image(str(png))
        img.width, img.height = img.width * 0.45, img.height * 0.45
        tk.add_image(img, "G3")

    # ---------------- Tablero Plan (solo si el ciclo trae plan: el plan anual se actualiza en los ciclos S2)
    tiene_plan = "plan.plan_plan" in T
    tpl = wb.create_sheet("Tablero_Plan")
    if not tiene_plan:
        tpl["A1"] = "Plan de compras: no se recalcula en este ciclo (el plan anual se actualiza en los ciclos S2)."
        tpl["A1"].font = TITULO
    else:
        _tablero_plan(tpl, T)

    # ---------------- Detalle
    p = T["desc.pareto_proveedores"].drop(columns=["ruc"], errors="ignore")
    hoja_datos(wb, "Pareto_ABC", p, "Hilo 3 – analizar.py", soles=("gasto",), pct=("pct_gasto", "pct_acumulado"))
    hoja_datos(wb, "Frecuencia", T["desc.frecuencia_proveedores"], "Hilo 3",
               soles=("gasto", "ticket_promedio_oc", "ticket_mediano_oc"))
    hoja_datos(wb, "Gasto_categoria", T["desc.gasto_categoria"], "Hilo 3", soles=("gasto",), pct=("pct_gasto",))
    hoja_datos(wb, "Gasto_subcategoria", T["desc.gasto_subcategoria"], "Hilo 3", soles=("gasto",), pct=("pct_gasto",))
    hoja_datos(wb, "Cruce_comprador_categoria", T["desc.especializacion_comprador"], "Hilo 3 (compradores con código)",
               soles=("gasto",), pct=("pct_categoria_principal",))
    hoja_datos(wb, "Concentracion_comprador", T["desc.concentracion_comprador"], "Hilo 3", soles=("gasto",),
               pct=("pct_comprador_principal",))
    hoja_datos(wb, "Sedes", T["desc.gasto_sede"], "Hilo 3", soles=("gasto",), pct=("pct_gasto",))
    hoja_datos(wb, "Alertas_resumen", T["alert.resumen"], "Hilo 4 – detalle completo en alertas.xlsx", soles=("monto_soles",))
    hoja_datos(wb, "Kraljic_unidades", T["kraljic.kraljic_unidades"].drop(columns=["justificacion_ia"], errors="ignore"),
               "Hilo 5 (provisional)", soles=("gasto",), pct=("pct_gasto", "pct_acumulado"))
    if tiene_plan:
        hoja_datos(wb, "Plan_anual", T["plan.plan_plan"], "Hilo 6 (ejercicio metodológico)",
                   soles=("linea_base", "rango_min", "rango_max", "tendencia_comparacion", "oc_tipica"))
        hoja_datos(wb, "Calendario", T["plan.plan_calendario"], "Hilo 6", soles=("linea_base",))
    hoja_datos(wb, "Control_calidad", calidad, "Hilo 1 – validaciones de la exportación")
    ctrl_df = T["desc.control"]
    ws_c = wb["Control_calidad"]
    escribir_tabla(ws_c, ctrl_df, fila=6 + len(calidad) + 2, soles=("valor_soles",), filtro=False)
    glos = pd.DataFrame([
        ("ABC", "Clase A: proveedores que suman el 80 % del gasto; B: hasta 95 %; C: el resto."),
        ("HHI", "Índice de concentración: suma de las participaciones al cuadrado × 10 000. > 1 800 = alta dependencia."),
        ("Kraljic", "Matriz impacto (gasto) × riesgo de suministro: estratégico, apalancamiento, cuello de botella, no crítico."),
        ("Fraccionamiento", "OC del mismo proveedor, sede y subcategoría en 60 días que, sumadas, superan un límite de aprobación."),
        ("Autoaprobación", "OC en la que el comprador y el aprobador registrado son la misma persona."),
        ("Adicionales de obra", "Trabajos no previstos en el contrato original de una obra."),
        ("Línea base", "Estimación inicial del plan de compras; se ajusta con el plan de obras y el presupuesto."),
        ("Soles nominales", "Montos sin ajustar por inflación."),
    ], columns=["Término", "Definición"])
    hoja_datos(wb, "Glosario", glos, "Proyecto análisis de gastos", ancho_max=110)
    for w in wb.worksheets:
        w.sheet_view.showGridLines = w.title.startswith("Tablero") is False
        w.page_setup.orientation = "landscape"
        w.page_setup.fitToWidth = 1
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return ruta


# --------------------------------------------------------------------------- modelo en estrella

def modelo_estrella(t: pd.DataFrame, carpeta: Path, kraljic_u: pd.DataFrame | None = None) -> dict[str, pd.DataFrame]:
    """Hechos (líneas de gasto) + dimensiones, con códigos en lugar de nombres de personas."""
    carpeta.mkdir(parents=True, exist_ok=True)
    fecha = t["periodo_factura"].fillna(t["fecha_oc"]).dt.to_period("M").dt.to_timestamp()
    hechos = pd.DataFrame({"id_linea": t["id_linea"], "fecha": fecha, "oc": t["oc"], "proveedor_id": t["proveedor_id"],
                           "comprador_codigo": t["comprador_codigo"], "aprobador_codigo": t["aprobador_codigo"],
                           "sede": t["sede"].fillna("(sin dato)"), "departamento": t["departamento"].fillna("(sin dato)"),
                           "categoria": t["categoria"], "subcategoria": t["subcategoria"], "producto": t["producto"],
                           "monto_gasto": t["monto_gasto"]})
    dims = {
        "dim_fecha": pd.DataFrame({"fecha": sorted(hechos["fecha"].dropna().unique())}).assign(
            anio=lambda d: d["fecha"].dt.year, mes=lambda d: d["fecha"].dt.month,
            semestre=lambda d: d["fecha"].dt.year.astype(str) + "-S" + ((d["fecha"].dt.month > 6) + 1).astype(str)),
        "dim_proveedor": t.groupby("proveedor_id").agg(proveedor=("proveedor", "first")).reset_index(),
        "dim_comprador": pd.DataFrame({"comprador_codigo": sorted(set(hechos["comprador_codigo"].dropna()) |
                                                                  set(hechos["aprobador_codigo"].dropna()))}),
        "dim_categoria": hechos[["categoria", "subcategoria", "producto"]].drop_duplicates().reset_index(drop=True),
        "dim_sede": pd.DataFrame({"sede": sorted(hechos["sede"].unique())}),
        "dim_departamento": pd.DataFrame({"departamento": sorted(hechos["departamento"].unique())}),
    }
    tablas = {"hechos_gasto": hechos, **dims}
    for k, v in tablas.items():
        v.to_parquet(carpeta / f"{k}.parquet", index=False)
    return tablas


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Excel de resultados con tablero (Hilo 7).")
    ap.add_argument("--descriptivo", required=True)
    ap.add_argument("--alertas", required=True)
    ap.add_argument("--kraljic", required=True)
    ap.add_argument("--plan", help="Carpeta del Hilo 6 (opcional: solo en ciclos S2)")
    ap.add_argument("--tabla", help="tabla_clasificada.parquet (para el modelo en estrella)")
    ap.add_argument("--calidad", help="reporte_calidad.xlsx del Hilo 1")
    ap.add_argument("--periodo", required=True)
    ap.add_argument("--salida", required=True)
    a = ap.parse_args(argv)
    T = cargar_tablas({"desc": a.descriptivo, "alert": a.alertas, "kraljic": a.kraljic, "plan": a.plan})
    T["_kraljic_png"] = str(Path(a.kraljic) / "graficos" / "kraljic_unidades.png")
    ruta = generar_resultados(T, leer_calidad(a.calidad), a.periodo, Path(a.salida) / f"resultados_{a.periodo}.xlsx")
    if a.tabla:
        modelo_estrella(pd.read_parquet(a.tabla), Path(a.salida) / "tablero")
    print("Generado:", ruta)
    return 0


if __name__ == "__main__":
    sys.exit(main())
